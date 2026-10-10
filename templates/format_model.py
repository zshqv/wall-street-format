"""
Wall Street Financial Model Formatter
Applies institutional investment banking formatting standards to Excel workbooks.

Usage:
    python format_model.py input.xlsx --output formatted.xlsx
    python format_model.py input.xlsx --sheet "Income Statement"
"""

import argparse
import re
import sys
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    NamedStyle,
    PatternFill,
    Side,
    numbers,
)
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# IB Color Constants
# ---------------------------------------------------------------------------
COLOR_INPUT = "0000FF"       # Blue  – hardcoded inputs
COLOR_FORMULA = "000000"     # Black – formulas / calculations
COLOR_EXT_LINK = "008000"    # Green – external workbook links
COLOR_WARNING = "FF0000"     # Red   – warnings / errors
COLOR_XSHEET = "800000"      # Maroon – cross-sheet references
COLOR_HEADER_FILL = "002060" # Dark navy header background
COLOR_HEADER_FONT = "FFFFFF" # White header text
COLOR_GRIDLINE = "D3D3D3"    # Light gray border

# ---------------------------------------------------------------------------
# Standard Font
# ---------------------------------------------------------------------------
FONT_NAME = "Arial"
FONT_SIZE = 10

# ---------------------------------------------------------------------------
# Number Formats
# ---------------------------------------------------------------------------
FMT_CURRENCY = '$#,##0_);($#,##0)'
FMT_CURRENCY_1D = '$#,##0.0_);($#,##0.0)'
FMT_PERCENT = '0.0%'
FMT_PERCENT_2D = '0.00%'
FMT_MULTIPLE = '0.0"x"'
FMT_MULTIPLE_2D = '0.00"x"'
FMT_DASH_ZERO = '#,##0_);(#,##0);"-"'
FMT_NUMBER = '#,##0_);(#,##0)'

# ---------------------------------------------------------------------------
# Border Styles
# ---------------------------------------------------------------------------
THIN_BORDER = Side(style="thin", color=COLOR_GRIDLINE)
THIN_DARK = Side(style="thin", color="000000")
DOUBLE_BORDER = Side(style="double", color="000000")
NO_BORDER = Side(style=None)

GRID_BORDER = Border(
    left=THIN_BORDER, right=THIN_BORDER,
    top=THIN_BORDER, bottom=THIN_BORDER,
)
SUBTOTAL_BORDER = Border(
    top=THIN_DARK, bottom=THIN_DARK,
    left=THIN_BORDER, right=THIN_BORDER,
)
TOTAL_BORDER = Border(
    top=THIN_DARK, bottom=DOUBLE_BORDER,
    left=THIN_BORDER, right=THIN_BORDER,
)

# ---------------------------------------------------------------------------
# Fills
# ---------------------------------------------------------------------------
HEADER_FILL = PatternFill(start_color=COLOR_HEADER_FILL,
                          end_color=COLOR_HEADER_FILL,
                          fill_type="solid")


def _is_total_row(row_cells):
    """Heuristic: a row is a total/subtotal if its first text cell contains
    keywords like Total, Subtotal, Net, EBITDA, etc."""
    for cell in row_cells:
        if cell.value and isinstance(cell.value, str):
            text = cell.value.strip().lower()
            if re.match(
                r"(total|subtotal|net |ebitda|ebit |gross profit|"
                r"free cash flow|net income|operating income)",
                text,
            ):
                return True
            break
    return False


def _is_header_row(row_cells):
    """Heuristic: a row is a header if every non-empty cell is a string
    (no numbers) and at least two cells have values."""
    values = [c.value for c in row_cells if c.value is not None]
    if len(values) < 2:
        return False
    return all(isinstance(v, str) for v in values)


def _classify_cell_color(cell, wb):
    """Determine the IB font color for a cell based on its content type."""
    if cell.value is None:
        return COLOR_FORMULA

    if cell.data_type == "f" or (
        isinstance(cell.value, str) and str(cell.value).startswith("=")
    ):
        formula = str(cell.value)
        if re.search(r"\[.*?\]", formula):
            return COLOR_EXT_LINK
        if "!" in formula:
            return COLOR_XSHEET
        return COLOR_FORMULA

    if isinstance(cell.value, (int, float)):
        return COLOR_INPUT

    return COLOR_FORMULA


def _guess_number_format(cell, col_idx, header_values):
    """Guess an appropriate IB number format based on header text or value."""
    if cell.value is None or isinstance(cell.value, str):
        return None

    header = None
    if header_values and col_idx < len(header_values):
        header = str(header_values[col_idx] or "").lower()

    if header:
        if any(k in header for k in ("margin", "growth", "yield", "%", "rate", "pct")):
            return FMT_PERCENT
        if any(k in header for k in ("multiple", "ev/", "p/e", "/x", "x ")):
            return FMT_MULTIPLE
        if any(k in header for k in ("$", "revenue", "ebitda", "income", "cash",
                                      "capex", "debt", "price", "cost", "expense")):
            return FMT_CURRENCY

    val = cell.value
    if isinstance(val, (int, float)):
        if 0 < abs(val) < 1.5:
            return FMT_PERCENT
        if abs(val) >= 1000:
            return FMT_CURRENCY

    return FMT_DASH_ZERO


def format_sheet(ws, wb):
    """Apply full Wall Street formatting to a single worksheet."""
    max_row = ws.max_row
    max_col = ws.max_column
    if max_row is None or max_col is None or max_row < 1:
        return

    header_values = []
    for col_idx in range(max_col):
        cell = ws.cell(row=1, column=col_idx + 1)
        header_values.append(cell.value)

    for row_idx in range(1, max_row + 1):
        row_cells = [ws.cell(row=row_idx, column=c + 1) for c in range(max_col)]
        is_header = (row_idx == 1) or _is_header_row(row_cells)
        is_total = _is_total_row(row_cells)

        for col_idx, cell in enumerate(row_cells):
            font_color = COLOR_HEADER_FONT if is_header else _classify_cell_color(cell, wb)

            cell.font = Font(
                name=FONT_NAME,
                size=FONT_SIZE,
                bold=is_header or is_total,
                color=font_color,
            )

            if is_header:
                cell.fill = HEADER_FILL
                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
            elif col_idx == 0:
                indent = 0
                if cell.value and isinstance(cell.value, str):
                    stripped = cell.value
                    leading = len(stripped) - len(stripped.lstrip())
                    indent = min(leading // 2, 4)
                cell.alignment = Alignment(
                    horizontal="left", vertical="center", indent=indent
                )
            else:
                cell.alignment = Alignment(
                    horizontal="right", vertical="center"
                )

            if is_total:
                cell.border = TOTAL_BORDER
            elif is_header:
                cell.border = Border(
                    bottom=THIN_DARK,
                    left=THIN_BORDER, right=THIN_BORDER,
                )
            else:
                cell.border = GRID_BORDER

            if not is_header and col_idx > 0:
                fmt = _guess_number_format(cell, col_idx, header_values)
                if fmt:
                    cell.number_format = fmt

    for col_idx in range(max_col):
        col_letter = get_column_letter(col_idx + 1)
        max_width = 10
        for row_idx in range(1, min(max_row + 1, 50)):
            cell = ws.cell(row=row_idx, column=col_idx + 1)
            if cell.value:
                max_width = max(max_width, min(len(str(cell.value)) + 2, 30))
        ws.column_dimensions[col_letter].width = max_width

    ws.sheet_view.showGridLines = False


def format_workbook(input_path, output_path=None, sheet_name=None):
    """Load a workbook, apply IB formatting, and save."""
    input_path = Path(input_path)
    if not input_path.exists():
        print(f"Error: file not found: {input_path}", file=sys.stderr)
        sys.exit(1)

    if output_path is None:
        output_path = input_path.with_stem(input_path.stem + "_formatted")
    else:
        output_path = Path(output_path)

    wb = load_workbook(str(input_path))

    sheets = [wb[sheet_name]] if sheet_name else wb.worksheets
    for ws in sheets:
        print(f"Formatting sheet: {ws.title}")
        format_sheet(ws, wb)

    wb.save(str(output_path))
    print(f"Saved formatted workbook to: {output_path}")


def main():
    parser = argparse.ArgumentParser(
        description="Apply Wall Street IB formatting to an Excel workbook."
    )
    parser.add_argument("input", help="Path to the input .xlsx file")
    parser.add_argument("--output", "-o", help="Path for the formatted output file")
    parser.add_argument("--sheet", "-s", help="Format only this sheet (by name)")
    args = parser.parse_args()

    format_workbook(args.input, args.output, args.sheet)


if __name__ == "__main__":
    main()
